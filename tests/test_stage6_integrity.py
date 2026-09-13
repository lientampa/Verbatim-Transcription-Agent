from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace
import json
import os
import docx
import pytest

from src.output_merger import OutputMerger
from src.output_renderer import OutputRenderer, format_seconds_to_srt_time
from src.renderer_validator import RendererValidator
from src.render_transaction import NAMES
from src.response_parser import TranscriptSegment

SOURCE = {"fingerprint": "a"*64, "duration_us": 3753250249}


def record(n, start, end, text="ờ ờ anh... anh nghĩ là, là cái này [không rõ]", stamp=None):
    seconds = int(start+1) if stamp is None else stamp
    return {"metric": {"block_id":f"BLOCK_{n:03d}","actual_start_offset":start,"actual_end_offset":end,
            "fidelity_decision":"ACCEPT","coverage":{"decision":"COVERAGE_PASS"},"segment_count":1},
            "segments":[{"source_index":1,"timestamp":f"{seconds//3600:02d}:{seconds%3600//60:02d}:{seconds%60:02d}","speaker":"Anh Nam","text":text}]}


def test_physical_sort_duplicates_and_same_timestamp_stability():
    a,b,c = record(1,0,600),record(2,600,900),record(3,900,1500)
    records = [c,a,b,deepcopy(b)]
    before = deepcopy(records)
    merger = OutputMerger.from_committed_records(records,SOURCE)
    assert len(merger.get_merged_segments()) == 3
    assert records == before
    a["segments"].append(dict(a["segments"][0], source_index=2,speaker="Chị Hoa"))
    a["metric"]["segment_count"] = 2
    merged = OutputMerger.from_committed_records([a],SOURCE).get_merged_segments()
    assert [s.speaker for s in merged] == ["Anh Nam","Chị Hoa"]


@pytest.mark.parametrize("kind,reason", [("payload","MERGE_CONFLICT"),("range","MULTIPLE_PHYSICAL"),("gap","MERGE_GAP"),("overlap","MERGE_OVERLAP"),("logical","ORDER_CONFLICT"),("rejected","REJECTED"),("timestamp","TIMESTAMP_ORDER")])
def test_merge_corruption_rejected(kind, reason):
    a,b=record(1,0,600),record(2,600,900)
    if kind=="payload":
        b=deepcopy(a); b["segments"][0]["text"]="conflict"
    if kind=="range": b["metric"]["block_id"]="BLOCK_001"
    if kind=="gap": b["metric"]["actual_start_offset"]=620
    if kind=="overlap": b["metric"]["actual_start_offset"]=590
    if kind=="logical": b["metric"]["block_id"]="BLOCK_003"
    if kind=="rejected": b["metric"]["coverage"]["decision"]="COVERAGE_RETRY"
    if kind=="timestamp": b["segments"][0]["timestamp"]="00:00:00"
    with pytest.raises(ValueError,match=reason):
        OutputMerger.from_committed_records([a,b],SOURCE)


def test_render_parity_verbatim_nonmutation_idempotent(tmp_path):
    segments=[TranscriptSegment(1,"Ờ... [không rõ] ờ ờ anh... anh nghĩ là, là cái này","00:03:14","Anh Nam"),
              TranscriptSegment(2,"source_index BLOCK_003 actual_model= là lời nói.","01:02:21","Chị Hoa")]
    before=deepcopy(segments)
    renderer=OutputRenderer(strict_speaker_format=True,source_end=3753.250249)
    renderer.render_all(segments,tmp_path)
    assert segments==before
    text=(tmp_path/"transcript.txt").read_text(encoding="utf-8")
    assert text.splitlines()[0]=="[00:03:14 - Anh Nam]: "+segments[0].text
    assert [p.text for p in docx.Document(tmp_path/"transcript.docx").paragraphs][1:]==text.splitlines()
    first={n:(tmp_path/n).read_bytes() for n in ("transcript.txt","subtitle.srt")}
    renderer.render_all(segments,tmp_path)
    assert first=={n:(tmp_path/n).read_bytes() for n in first}
    assert "01:02:21,000 --> 01:02:25,000" in (tmp_path/"subtitle.srt").read_text(encoding="utf-8")


@pytest.mark.parametrize("speaker",["SPEAKER_01","segment_81","BLOCK_003","actual_model=gemini-3.6-flash"])
def test_technical_speaker_rejected(tmp_path,speaker):
    with pytest.raises(ValueError):
        OutputRenderer(strict_speaker_format=True).render_all([TranscriptSegment(1,"x","00:00:01",speaker)],tmp_path)
    assert not (tmp_path/"transcript.txt").exists()


def test_srt_next_start_same_start_and_eof():
    segs=[TranscriptSegment(i,"x",ts,"Anh") for i,ts in enumerate(["00:00:10","00:00:14","00:00:18"],1)]
    renderer=OutputRenderer(strict_speaker_format=True,source_end=20)
    assert renderer.srt_times(segs)==[(10,14),(14,18),(18,20)]
    segs[1].timestamp="00:00:10"
    assert renderer.srt_times(segs)[0]==(10,10.1)
    assert segs[0].timestamp==segs[1].timestamp=="00:00:10"
    with pytest.raises(ValueError): OutputRenderer(srt_end_policy="MODEL_END")
    assert format_seconds_to_srt_time(59.9999)=="00:01:00,000"


@pytest.mark.parametrize("phase",["docx_render","docx_publish","parity"])
def test_atomic_old_set_survives_failure(tmp_path,monkeypatch,phase):
    renderer=OutputRenderer(strict_speaker_format=True)
    renderer.render_all([TranscriptSegment(1,"old","00:00:01","Anh")],tmp_path)
    original={n:(tmp_path/n).read_bytes() for n in NAMES}
    if phase=="docx_render":
        monkeypatch.setattr(renderer,"render_docx",lambda *a,**k: (_ for _ in ()).throw(OSError("locked")))
    elif phase=="parity":
        original_docx=renderer.render_docx
        def bad(segs,path):
            original_docx(segs,path)
            document=docx.Document(path); document.paragraphs[-1].text="lost segment"; document.save(path)
        monkeypatch.setattr(renderer,"render_docx",bad)
    else:
        replace=os.replace
        def fail(src,dst):
            if Path(dst)==tmp_path/"transcript.docx" and ".render-stage-" in str(src): raise OSError("locked")
            return replace(src,dst)
        monkeypatch.setattr(os,"replace",fail)
    with pytest.raises((ValueError,OSError)):
        renderer.render_all([TranscriptSegment(1,"new","00:00:01","Anh")],tmp_path)
    assert original=={n:(tmp_path/n).read_bytes() for n in NAMES}


def test_production_checkpoint_eof_render_six_blocks(tmp_path,monkeypatch):
    from tests.test_stage4_audio_retry import run_audio_case
    from src.checkpoint import CheckpointManager
    from src.source_identity import SourceIdentity
    from dataclasses import replace
    ends=[600,900,1500,2100,2700,3753.250249]
    def setup(config):
        config=replace(config,strict_speaker_format=True)
        monkeypatch.setattr("src.main.load_config",lambda **kw:config)
        manager=CheckpointManager(config.checkpoint_file_path)
        manager.create_new_job("source.wav",source_identity=SourceIdentity.capture(config.audio_dir/"source.wav",ends[-1]))
        start=0
        for n,end in enumerate(ends,1):
            r=record(n,start,end,stamp=3741 if n==6 else None)
            manager.commit_block(r["metric"]["block_id"],1,r["segments"],metric=dict(r["metric"],source_mode="AUDIO"))
            start=end
    result,_,uploads,calls,cp=run_audio_case(tmp_path,monkeypatch,setup=setup,fail_count=0,total_duration=ends[-1])
    assert result==0 and not uploads and not calls
    output=tmp_path/"output"
    assert len((output/"transcript.txt").read_text(encoding="utf-8").splitlines())==6
    assert len(docx.Document(output/"transcript.docx").paragraphs)==7
    assert "01:02:21,000 --> 01:02:25,000" in (output/"subtitle.srt").read_text(encoding="utf-8")
    assert json.loads((output/"render_manifest.json").read_text())["committed_block_count"]==6


def test_text_mode_remains_source_index_order():
    merger=OutputMerger([TranscriptSegment(2,"b"),TranscriptSegment(1,"a")])
    assert [s.text for s in merger.get_merged_segments()]==["a","b"]


def test_rejected_generations_never_rendered_on_resume(tmp_path,monkeypatch):
    from tests.test_stage4_audio_retry import run_audio_case
    from src.checkpoint import CheckpointManager
    from src.source_identity import SourceIdentity
    from src.gemini_client import GeminiClient
    from dataclasses import replace
    def setup(config):
        config=replace(config,next_target_policy="MAX_FIRST",max_duration_seconds=1200,initial_target_input_tokens=3000,strict_speaker_format=True)
        monkeypatch.setattr("src.main.load_config",lambda **kw:config)
        manager=CheckpointManager(config.checkpoint_file_path)
        manager.create_new_job("source.wav",source_identity=SourceIdentity.capture(config.audio_dir/"source.wav",1800))
        r=record(1,0,600,text="accepted first")
        manager.commit_block("BLOCK_001",1,r["segments"],metric=dict(r["metric"],source_mode="AUDIO"))
        original=GeminiClient.generate_transcription
        def generate(self,**kw):
            data=json.loads(original(self,**kw))
            start,end=kw["gemini_file"].bounds
            rejected=start==600 and end-start>300
            sec=int(start+1 if rejected else end-2)
            data["segments"][0].update(text="rejected generation" if rejected else "accepted retry",speaker="Anh Nam",timestamp=f"{sec//60:02d}:{sec%60:02d}")
            return json.dumps(data)
        monkeypatch.setattr(GeminiClient,"generate_transcription",generate)
        monkeypatch.setattr("src.coverage_validator.TailActivityAnalyzer.analyze",lambda self,path,offset,duration,threshold:duration)
    result,_,_,calls,cp=run_audio_case(tmp_path,monkeypatch,setup=setup,fail_count=0,total_duration=1800)
    assert result==0
    assert [c[0].bounds for c in calls[:3]]==[[600,1800],[600,1200],[600,900]]
    assert cp["block_metrics"][1]["actual_end_offset"]==900
    output=tmp_path/"output"
    for name in ("transcript.txt","subtitle.srt"):
        assert "rejected generation" not in (output/name).read_text(encoding="utf-8")
    assert all("rejected generation" not in p.text for p in docx.Document(output/"transcript.docx").paragraphs)


def test_pending_publication_is_rolled_back_after_process_loss(tmp_path,monkeypatch):
    import src.render_transaction as transaction
    renderer=OutputRenderer(strict_speaker_format=True)
    renderer.render_all([TranscriptSegment(1,"old","00:00:01","Anh")],tmp_path)
    old={name:(tmp_path/name).read_bytes() for name in NAMES}
    real_replace=os.replace
    def fail(src,dst):
        if Path(dst)==tmp_path/"transcript.docx" and ".render-stage-" in str(src):
            raise KeyboardInterrupt("process lost")
        return real_replace(src,dst)
    real_recover=transaction.recover_outputs
    def recovery(directory):
        if (directory/".render-transaction"/"pending.json").exists():
            raise KeyboardInterrupt("no cleanup on abrupt loss")
        return real_recover(directory)
    with monkeypatch.context() as patch:
        patch.setattr(os,"replace",fail)
        patch.setattr(transaction,"recover_outputs",recovery)
        with pytest.raises(KeyboardInterrupt):
            renderer.render_all([TranscriptSegment(1,"new","00:00:01","Anh")],tmp_path)
    real_recover(tmp_path)
    assert old=={name:(tmp_path/name).read_bytes() for name in NAMES}
